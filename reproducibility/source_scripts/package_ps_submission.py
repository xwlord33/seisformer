"""Refresh the local v7 manuscript/reproduction package after successful checks."""
from pathlib import Path
import hashlib
import json
import re
import shutil
import zipfile

PROJECT=Path(__file__).resolve().parents[1]
ROOT=PROJECT/'outputs/paper_ps'


def sha(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream,'sha256').hexdigest()


def run():
    check=json.loads((ROOT/'analysis/manuscript_checks.json').read_text(encoding='utf-8'))
    assert check['status']=='passed'
    for name,record in check['rendered_pdfs'].items():
        assert sha(ROOT/(name+'.pdf'))==record['sha256'], 'Re-run manuscript checks after changing PDFs.'
    alignment=json.loads((ROOT/'analysis/bilingual_alignment.json').read_text(encoding='utf-8'))
    # Alignment hashes the decoded text with normalized newlines; package
    # manifest entries below hash raw bytes, including Windows CRLF if present.
    assert hashlib.sha256((ROOT/'main.tex').read_text(encoding='utf-8').encode()).hexdigest()==alignment['english_master_sha256']
    snapshots=ROOT/'reproducibility/source_scripts'
    snapshots.mkdir(exist_ok=True)
    import docx
    from docx.shared import Pt
    import subprocess
    doc_h=docx.Document()
    doc_h.styles['Normal'].font.name='Times New Roman'
    doc_h.styles['Normal'].font.size=Pt(11)
    doc_h.add_heading('Highlights',level=1)
    for line in (ROOT/'HIGHLIGHTS.md').read_text(encoding='utf-8').splitlines():
        if line.startswith('- '):
            doc_h.add_paragraph(line[2:].strip(),style='List Bullet')
    doc_h.save(ROOT/'Highlights.docx')
    doc_c=docx.Document()
    doc_c.styles['Normal'].font.name='Times New Roman'
    doc_c.styles['Normal'].font.size=Pt(11)
    for block in (ROOT/'COVER_LETTER_CG.md').read_text(encoding='utf-8').strip().split('\n\n'):
        clean=block.replace('**','').replace('*','')
        p=doc_c.add_paragraph(clean)
        p.paragraph_format.space_after=Pt(8)
    doc_c.save(ROOT/'Cover_letter.docx')
    h_tex=(
        "\\documentclass[11pt,a4paper]{article}\n\\usepackage[margin=25mm]{geometry}\n"
        "\\usepackage[T1]{fontenc}\n\\usepackage[utf8]{inputenc}\n\\usepackage{lmodern,microtype}\n"
        "\\pagestyle{empty}\n\\begin{document}\n\\section*{Highlights}\n\\begin{itemize}\n"
        "\\item Aligns Morlet and STFT features to 375 tokens with only 5,856 added parameters.\n"
        "\\item Paired 262k seeds reduce P/S arrival MAE from 0.551/0.346~s to 0.473/0.313~s.\n"
        "\\item Cuts P/S MAE by 43.4\\%/21.6\\% vs.\\ TwoStream-EQT with 31.1\\% fewer parameters.\n"
        "\\item Primary 1M model reaches 0.266/0.216~s P/S MAE and 4.18~ms GPU window latency.\n"
        "\\item Source-specific conformal calibration cuts 90\\% coverage deviation to 0.017.\n"
        "\\end{itemize}\n\\end{document}\n"
    )
    (ROOT/'highlights.tex').write_text(h_tex,encoding='utf-8')
    c_tex=(
        "\\documentclass[11pt,a4paper]{article}\n\\usepackage[margin=25mm]{geometry}\n"
        "\\usepackage[T1]{fontenc}\n\\usepackage[utf8]{inputenc}\n\\usepackage{lmodern,microtype,hyperref}\n"
        "\\hypersetup{colorlinks=true,urlcolor=blue}\n\\setlength{\\parskip}{6pt}\n\\setlength{\\parindent}{0pt}\n"
        "\\pagestyle{empty}\n\\begin{document}\nDear Editors,\n\n"
        "Please consider our manuscript, \\textbf{``XFNet: compact time--frequency fusion for P- and S-wave arrival picking across multi-source seismic archives,''} for publication in \\textit{Computers \\& Geosciences}.\n\n"
        "The study addresses the computational problem of adding time--frequency representations to a single-station seismic phase picker while keeping parameter growth small and measuring the resulting accuracy--cost trade-off. XFNet (Cross-domain Time--Frequency Network) aligns fixed Morlet-filterbank and short-time Fourier views with a 375-token convolutional representation of a 60-s three-component waveform. The spectral paths add 5,856 trainable parameters (1.38\\%) to the temporal control, giving 428,774 parameters in total.\n\n"
        "Three paired fine-tuning seeds using identical training records and a shared pilot initialization reduce mean P/S arrival errors from 0.551/0.346~s to 0.473/0.313~s. A separate matched-budget comparison with the public cross-attention configuration of TwoStream-EQT uses three random-initialization seeds on the same 262,144 training records; XFNet reduces mean P/S MAE from 0.751/0.376~s to 0.425/0.295~s with 31.1\\% fewer parameters and lower single-window latency, while the reference achieves higher batch-64 throughput. A larger-data continuation achieves P/S MAE 0.266/0.216~s (81.9\\%/72.9\\% of picks within $\\pm0.1$~s; signed mean biases $+0.042$/$+0.013$~s) and detection F1 0.9988 on the stated multi-source test population. Direct checkpoint timing gives 4.18~ms per 60-s input window on an RTX~5070~Ti. We also examine source-dependent arrival intervals, linking representation quality to the observations supplied to later processing. Held-out STEAD, TXED and PNW waveform examples connect these results to measured arrivals.\n\n"
        "The manuscript is original, is not under consideration elsewhere, and all authors have approved the final submission. The complete source code, primary model weights, calibration manifests and reproduction scripts are openly available under the MIT License at \\url{https://github.com/xwlord33/xfnet}. We believe the combination of compact algorithm design, controlled ablation, matched recent-model benchmarking and reproducible seismic-data analysis is well suited to the readership of \\textit{Computers \\& Geosciences}.\n\n"
        "Sincerely,\n\n"
        "\\textbf{Wei Xiang} (Corresponding author; on behalf of all authors: Wei Xiang and Jianying Mou)\\\\\n"
        "1.\\ Guangxi Zhuang Autonomous Region Earthquake Agency, No.\\ 2 Leitan Road, Liangqing District, Nanning 530022, China\\\\\n"
        "2.\\ Beihai Weizhou Island Marine Geophysics Observation and Research Station of Guangxi, Weizhou Island Earthquake Monitoring Center Station, No.\\ 1 Planning Road, Weizhou Town, Haicheng District, Beihai 536000, China\\\\\n"
        "Email: \\texttt{xiangwei@gxearthquake.cn} $|$ ORCID: \\texttt{0009-0001-4572-3507}\\\\\n"
        "Co-author: Jianying Mou (Guangxi Zhuang Autonomous Region Earthquake Agency; Email: \\texttt{645335110@qq.com})\n"
        "\\end{document}\n"
    )
    (ROOT/'cover_letter.tex').write_text(c_tex,encoding='utf-8')
    for src,job in [('highlights.tex','Highlights'),('cover_letter.tex','Cover_letter')]:
        subprocess.run(['pdflatex','-interaction=nonstopmode','-halt-on-error','-jobname='+job,src],cwd=ROOT,check=True,capture_output=True)
    for name in ('benchmark_inference','build_paper_ps_tex','build_ps_supplement',
                 'check_ps_submission','compile_ps_submission','export_ps_manuscript_markdown',
                 'make_ps_submission_figures','make_ps_heldout_waveforms','prepare_ps_submission_evidence',
                 'build_ps_chinese_review','search_ps_journal_literature','package_ps_submission',
                 'recent_dual_core','prepare_recent_dual_comparison','preflight_recent_dual_comparison',
                 'train_recent_dual_comparison','run_recent_dual_comparison','analyze_recent_dual_comparison',
                 'export_recent_dual_paper_table','export_recent_dual_repro','integrate_recent_dual_manuscript',
                 'rebase_ps_review_translation','diagnose_recent_dual_nan','probe_recent_dual_precision'):
        shutil.copy2(PROJECT/'scripts'/(name+'.py'),snapshots/(name+'.py'))
    files=[]
    for name in ('main.tex','supplement.tex','chinese_review.tex','refs.bib',
                 'XFNet_manuscript.pdf','XFNet_supplement.pdf','中文核对版.pdf',
                 'Highlights.docx','Highlights.pdf','Cover_letter.docx','Cover_letter.pdf',
                 'MANUSCRIPT_EN.md','HIGHLIGHTS.md','中文核对版.md','中英逐段对照.md','README.md',
                 'COVER_LETTER_GJI.md','COVER_LETTER_CG.md','修改说明_GJI.md',
                 '投稿信息待填写.md','选刊与论文定位.md',
                 '论文模式修改对照_20261022.md','近期同类论文对照与投稿判断_20261022.md',
                 '近期双域公平对照结果.md',
                 'analysis/manuscript_checks.json','analysis/primary_runtime.json',
                 'analysis/references/crossref_metadata.json','analysis/submission_evidence.json',
                 'analysis/heldout_waveform_selection.json','analysis/heldout_waveform_data.npz',
                 'analysis/review_units.json','analysis/review_translation_zh.json',
                 'analysis/bilingual_alignment.json','analysis/build_logs/compilation.json',
                 'analysis/review_labels_zh.json',
                 'analysis/recent_dual_table.tex','analysis/recent_dual_supplement.tex','analysis/recent_dual_export.json'):
        files.append(ROOT/name)
    main=(ROOT/'main.tex').read_text(encoding='utf-8')
    figure_files=[]
    for figure in re.findall(r'\\includegraphics\[[^]]*\]\{([^}]+)\}',main):
        figure_files.extend((ROOT/figure).with_suffix(ext) for ext in ('.pdf','.svg','.png'))
    files.extend(figure_files)
    files.extend(p for p in (ROOT/'reproducibility').rglob('*') if p.is_file() and '__pycache__' not in p.parts
                 and p.name!='revise_ps_competition_scope.py')
    files.extend(p for p in (ROOT/'analysis/journal_recent_20261022').iterdir()
                 if p.suffix in ('.json','.md','.txt'))
    files=sorted(set(files),key=lambda p:p.relative_to(ROOT).as_posix())
    records=[dict(path=p.relative_to(ROOT).as_posix(),bytes=p.stat().st_size,sha256=sha(p)) for p in files]
    manifest=dict(version_date='2026-10-24',manuscript_revision='cg-ready',
                  target_journal='Computers & Geosciences',
                  alternative_journal='Geophysical Journal International',
                  purpose='Manuscript and reproduction materials',
                  files=records,primary_model_sha256=check['primary_weight_sha256'])
    manifest_path=ROOT/'XFNet_source_manifest.json'
    manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    target=ROOT/'XFNet_submission_materials.zip'
    temporary=target.with_suffix('.tmp.zip')
    with zipfile.ZipFile(temporary,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as archive:
        for p in files+[manifest_path]: archive.write(p,p.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(temporary) as archive:
        assert archive.testzip() is None
        for record in records:
            assert hashlib.sha256(archive.read(record['path'])).hexdigest()==record['sha256']
    temporary.replace(target)
    cg_files=sorted(set([
        ROOT/'main.tex', ROOT/'supplement.tex', ROOT/'refs.bib',
        ROOT/'XFNet_manuscript.pdf', ROOT/'XFNet_supplement.pdf',
        ROOT/'Highlights.docx', ROOT/'Highlights.pdf',
        ROOT/'Cover_letter.docx', ROOT/'Cover_letter.pdf',
        ROOT/'HIGHLIGHTS.md', ROOT/'COVER_LETTER_CG.md',
    ] + figure_files), key=lambda p:p.relative_to(ROOT).as_posix())
    cg_target=ROOT/'XFNet_CG_submission_upload.zip'
    cg_tmp=cg_target.with_suffix('.tmp.zip')
    with zipfile.ZipFile(cg_tmp,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as archive:
        for p in cg_files: archive.write(p,p.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(cg_tmp) as archive:
        assert archive.testzip() is None
    cg_tmp.replace(cg_target)
    print(json.dumps(dict(package=str(target),cg_upload_package=str(cg_target),files=len(records)+1,cg_upload_files=len(cg_files),MiB=round(target.stat().st_size/2**20,2),
                          all_archive_hashes_verified=True),ensure_ascii=False,indent=2))


if __name__=='__main__':run()
